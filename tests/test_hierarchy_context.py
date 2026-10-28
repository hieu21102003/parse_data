"""Tests for catalog hierarchy extraction (parent levels + STT).

The spec: a "Danh mục tiêu chuẩn" groups codes under Roman-numeral
groups (level 1: I, II, ...) and dotted subsections (level 2/3:
I.1, I.2, I.2.1, or the Arabic form "1.2."), with a TT (STT)
column on code rows. Every extracted record must carry the
nearest headings above it as "parent_level_1" / "parent_level_2"
and the row's "stt" -- never invented, None when the source has
no such context. The code itself is "corrected_code" and the
title is "standard_name", as in the rest of the pipeline.

Fixtures mirror the user's real catalog file: a title row, a
TT / Ký hiệu / Tên header, heading rows with the marker in the
first cell and the title in the next, numbered sub-headings in
the code column, and STT values that restart per section.
"""

try:
    from ingestion.standard_code_extractor import (
        _heading_from_row,
        _HierarchyContext,
        extract_standard_codes,
        parse_heading_text,
    )
except ModuleNotFoundError:  # pragma: no cover - repo-root PYTHONPATH
    from src.ingestion.standard_code_extractor import (  # type: ignore
        _heading_from_row,
        _HierarchyContext,
        extract_standard_codes,
        parse_heading_text,
    )

try:
    from pipeline import process_document, process_text
except ModuleNotFoundError:  # pragma: no cover - repo-root PYTHONPATH
    from src.pipeline import process_document, process_text  # type: ignore


def test_parse_heading_text_levels():
    assert parse_heading_text("I. Nhóm công trình đường sắt; đepot đường sắt") == (
        1, "I. Nhóm công trình đường sắt; đepot đường sắt",
    )
    assert parse_heading_text("II. Nhóm thông tin") == (1, "II. Nhóm thông tin")
    assert parse_heading_text("I.2.1. Chung, tuyến, kết cấu tầng trên") == (
        2, "I.2.1. Chung, tuyến, kết cấu tầng trên",
    )
    assert parse_heading_text("I.1. Quy chuẩn kỹ thuật") == (
        2, "I.1. Quy chuẩn kỹ thuật",
    )
    assert parse_heading_text("1.2. Tiêu chuẩn kỹ thuật") == (
        2, "1.2. Tiêu chuẩn kỹ thuật",
    )
    # Bare numbered lines are NOT headings in free text: too often
    # an ordinary numbered sentence.
    assert parse_heading_text("2. Thiết kế") is None
    # ... but they are headings in table rows (see _heading_from_row).
    assert parse_heading_text("2. Thiết kế", allow_bare_numbered=True) == (
        2, "2. Thiết kế",
    )


def test_parse_heading_text_rejects_non_headings():
    assert parse_heading_text("TT") is None
    assert parse_heading_text("Ký hiệu") is None
    assert parse_heading_text("TCVN 1234:2020") is None
    assert parse_heading_text("Phụ lục 4: Thống kê các tiêu chuẩn") is None
    assert parse_heading_text("") is None
    assert parse_heading_text("2020. Something") is None


def test_heading_from_row_forms():
    assert _heading_from_row(["I", "Nhóm công trình đường sắt"]) == (
        1, "I. Nhóm công trình đường sắt",
    )
    assert _heading_from_row(["I.1", "Quy chuẩn kỹ thuật quốc gia"]) == (
        2, "I.1. Quy chuẩn kỹ thuật quốc gia",
    )
    assert _heading_from_row(["I.1.", "Quy chuẩn kỹ thuật", "", "", "Không có"]) == (
        2, "I.1. Quy chuẩn kỹ thuật",
    )
    assert _heading_from_row(["1.2.", "Tiêu chuẩn kỹ thuật"]) == (
        2, "1.2. Tiêu chuẩn kỹ thuật",
    )
    assert _heading_from_row(["I. Nhóm công trình đường sắt"]) == (
        1, "I. Nhóm công trình đường sắt",
    )
    assert _heading_from_row(["", "2. Thiết kế"]) == (2, "2. Thiết kế")
    # A row with a code is never a heading row.
    assert _heading_from_row(["1", "JIS E 1102 - 2001", "Fish plates"]) is None
    assert _heading_from_row(["I", "EN 1993-1-1:2005"]) is None
    assert _heading_from_row(["TT", "Ký hiệu", "Tên tiêu chuẩn"]) is None


def test_hierarchy_context_reset_on_new_level_1():
    ctx = _HierarchyContext()
    ctx.update((1, "I. Nhóm A"))
    ctx.update((2, "I.2. Tiêu chuẩn kỹ thuật"))
    assert ctx.snapshot() == {
        "parent_level_1": "I. Nhóm A",
        "parent_level_2": "I.2. Tiêu chuẩn kỹ thuật",
    }
    ctx.update((1, "II. Nhóm B"))
    assert ctx.snapshot() == {"parent_level_1": "II. Nhóm B", "parent_level_2": None}
    ctx.update((2, "2. Thiết kế"))
    assert ctx.snapshot()["parent_level_2"] == "2. Thiết kế"


def _write_excel(tmp_path, sheets):
    import pandas as pd

    path = tmp_path / "catalog.xlsx"
    with pd.ExcelWriter(path) as writer:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=name, index=False, header=False)
    return str(path)


_CATALOG_SHEET_1 = [
    ["Phụ lục 4: Thống kê các tiêu chuẩn nước ngoài"],
    ["TT", "Ký hiệu", "Tên tiêu chuẩn"],
    ["I", "Nhóm công trình đường sắt; đepot đường sắt"],
    ["I.1", "Quy chuẩn kỹ thuật quốc gia"],
    ["I.2", "Tiêu chuẩn kỹ thuật"],
    ["", "2. Thiết kế"],
    ["1", "BS EN 1990:2002", "Cơ sở thiết kế kết cấu"],
    ["2", "BS EN 1991-1-1:2002", "Tác động lên kết cấu"],
    ["II", "Nhóm thông tin – tín hiệu và điều khiển"],
    ["1", "CISPR 22:2008", "Thiết bị công nghệ thông tin"],
]


def test_excel_hierarchy_and_stt(tmp_path):
    path = _write_excel(tmp_path, {"TH": _CATALOG_SHEET_1})
    records = process_document(path, use_ai_fallback=False)
    by_code = {r["corrected_code"]: r for r in records}

    first = by_code["BS EN 1990:2002"]
    assert first["parent_level_1"] == "I. Nhóm công trình đường sắt; đepot đường sắt"
    assert first["parent_level_2"] == "2. Thiết kế"
    assert first["stt"] == "1"
    assert first["standard_name"] == "Cơ sở thiết kế kết cấu"

    second = by_code["BS EN 1991-1-1:2002"]
    assert second["parent_level_2"] == "2. Thiết kế"
    assert second["stt"] == "2"

    # A new Roman group clears the subsection: no level-2 parent.
    third = by_code["CISPR 22:2008"]
    assert third["parent_level_1"] == "II. Nhóm thông tin – tín hiệu và điều khiển"
    assert third["parent_level_2"] is None
    assert third["stt"] == "1"


def test_excel_expansion_keeps_parents_and_stt(tmp_path):
    rows = [
        ["TT", "Ký hiệu", "Tên tiêu chuẩn"],
        ["I", "Nhóm công trình đường sắt"],
        ["I.2", "Tiêu chuẩn kỹ thuật"],
        ["1", "BS EN 13146-1 đến 10:2019", "Railway applications"],
    ]
    path = _write_excel(tmp_path, {"TH": rows})
    records = process_document(path, use_ai_fallback=False)
    assert len(records) == 10
    for record in records:
        assert record["parent_level_1"] == "I. Nhóm công trình đường sắt"
        assert record["parent_level_2"] == "I.2. Tiêu chuẩn kỹ thuật"
        assert record["stt"] == "1"


def test_excel_sheet_change_resets_hierarchy(tmp_path):
    sheet_2 = [
        ["TT", "Ký hiệu", "Tên tiêu chuẩn"],
        ["1", "GB 50157-2003", "Quy phạm thiết kế Metro"],
    ]
    path = _write_excel(tmp_path, {"TH": _CATALOG_SHEET_1, "Other": sheet_2})
    records = process_document(path, use_ai_fallback=False)
    by_code = {r["corrected_code"]: r for r in records}
    # Canonical year separator: "GB 50157-2003" -> "GB 50157:2003".
    assert by_code["GB 50157:2003"]["parent_level_1"] is None
    assert by_code["GB 50157:2003"]["parent_level_2"] is None
    assert by_code["GB 50157:2003"]["stt"] == "1"


def test_docx_heading_paragraph_governs_following_table(tmp_path):
    from docx import Document

    path = tmp_path / "catalog.docx"
    doc = Document()
    doc.add_paragraph("I. Nhóm công trình đường sắt; đepot đường sắt")
    doc.add_paragraph("I.2.1. Chung, tuyến, kết cấu tầng trên, nền, vật tư đường sắt")
    table = doc.add_table(rows=3, cols=3)
    headers = ["TT", "Ký hiệu", "Tên tiêu chuẩn"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "GB/T 44413-2024"
    table.rows[1].cells[2].text = "Tên tiêu chuẩn mẫu"
    table.rows[2].cells[0].text = "2"
    table.rows[2].cells[1].text = "EN 1993-1-1:2005"
    table.rows[2].cells[2].text = "Thiết kế kết cấu thép"
    doc.save(str(path))

    records = process_document(str(path), use_ai_fallback=False)
    by_code = {r["corrected_code"]: r for r in records}
    # The year separator is canonicalised ("-2024" -> ":2024") but
    # the "GB/T" designation itself is preserved, not reduced to "GB".
    first = by_code["GB/T 44413:2024"]
    assert first["raw_text"] == "GB/T 44413-2024"
    assert first["parent_level_1"] == "I. Nhóm công trình đường sắt; đepot đường sắt"
    assert first["parent_level_2"] == "I.2.1. Chung, tuyến, kết cấu tầng trên, nền, vật tư đường sắt"
    assert first["stt"] == "1"
    assert by_code["EN 1993-1-1:2005"]["stt"] == "2"


def test_docx_heading_row_inside_table(tmp_path):
    from docx import Document

    path = tmp_path / "catalog2.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "TT"
    table.rows[0].cells[1].text = "Ký hiệu"
    table.rows[1].cells[0].text = "I"
    table.rows[1].cells[1].text = "Nhóm công trình đường sắt"
    table.rows[2].cells[0].text = "I.1"
    table.rows[2].cells[1].text = "Quy chuẩn kỹ thuật"
    table.rows[3].cells[0].text = "1"
    table.rows[3].cells[1].text = "JIS E 1102-2001"
    doc.save(str(path))

    records = process_document(str(path), use_ai_fallback=False)
    assert len(records) == 1
    assert records[0]["parent_level_1"] == "I. Nhóm công trình đường sắt"
    assert records[0]["parent_level_2"] == "I.1. Quy chuẩn kỹ thuật"
    assert records[0]["stt"] == "1"


def test_text_lines_carry_headings_without_changing_matches():
    text = (
        "I. Nhóm công trình đường sắt\n"
        "I.2. Tiêu chuẩn kỹ thuật\n"
        "BS EN 1990:2002, BS EN 1991-1-1:2002\n"
        "II. Nhóm thông tin\n"
        "CISPR 22:2008"
    )
    records = process_text(text, use_ai_fallback=False)
    by_code = {r["corrected_code"]: r for r in records}
    assert by_code["BS EN 1990:2002"]["parent_level_1"] == "I. Nhóm công trình đường sắt"
    assert by_code["BS EN 1990:2002"]["parent_level_2"] == "I.2. Tiêu chuẩn kỹ thuật"
    assert by_code["CISPR 22:2008"]["parent_level_1"] == "II. Nhóm thông tin"
    assert by_code["CISPR 22:2008"]["parent_level_2"] is None
    # Plain text has no STT column.
    assert all(r["stt"] is None for r in records)


def test_no_heading_means_none_parents():
    records = process_text("BS EN 1990:2002", use_ai_fallback=False)
    assert len(records) == 1
    assert records[0]["parent_level_1"] is None
    assert records[0]["parent_level_2"] is None
    assert records[0]["stt"] is None


def test_extractor_records_carry_context_fields(tmp_path):
    path = _write_excel(tmp_path, {"TH": _CATALOG_SHEET_1})
    items = extract_standard_codes(path, use_ai_fallback=False)
    assert items, "expected extractor records"
    for item in items:
        assert "parent_level_1" in item
        assert "parent_level_2" in item
        assert "stt" in item
